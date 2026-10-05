from docx import Document
from pathlib import Path

kb_dir = Path(__file__).parents[1] / 'knowledge_base'
kb_dir.mkdir(parents=True, exist_ok=True)
doc_path = kb_dir / 'company_policy.docx'

doc = Document()
doc.add_heading('Company Policy — Customer Support', level=1)
doc.add_paragraph(
    'Purpose: This document defines the company policy for handling customer support requests, '
    'response SLAs, escalation procedures, and data handling guidelines.'
)

doc.add_heading('1. Scope', level=2)
doc.add_paragraph('This policy applies to all customer-facing support channels including chat, email, and phone.')

doc.add_heading('2. Response Times', level=2)
doc.add_paragraph('Priority 1 (Critical): initial response within 1 hour.\n'
                  'Priority 2 (High): initial response within 4 hours.\n'
                  'Priority 3 (Normal): initial response within 24 hours.')

doc.add_heading('3. Escalation', level=2)
doc.add_paragraph('If an issue is unresolved after 48 hours, escalate to Level 2 Support and notify the support manager.')

doc.add_heading('4. Data Handling', level=2)
doc.add_paragraph('Do not request or store sensitive personal data in chat logs. Use secure ticketing system for PII.')

doc.add_paragraph('\nLast updated: 2026-08-28')

doc.save(doc_path)
print(f'Wrote sample docx to: {doc_path}')
